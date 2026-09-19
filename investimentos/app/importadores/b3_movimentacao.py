"""Importador do Excel de Movimentacao da Area do Investidor da B3.

O arquivo sai em Relatorios > Movimentacao > Exportar para Excel. As colunas
sao, nessa ordem:

    Entrada/Saida | Data | Movimentacao | Produto | Instituicao | Quantidade |
    Preco unitario | Valor da Operacao

Esse arquivo e a fonte mais confiavel de proventos que existe para carteira
brasileira: cada dividendo, JCP, rendimento e amortizacao aparece com data de
credito e valor exato, ja liquido do que foi retido.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, BinaryIO, Iterable, Optional

from openpyxl import load_workbook
from sqlmodel import Session

from app.db import obter_ou_criar_ativo, obter_ou_criar_conta, registrar_movimentacao
from app.models import Fonte, Instituicao, Movimentacao, TipoMovimento

#: Texto da coluna "Movimentacao" mapeado para o tipo interno. As chaves estao
#: normalizadas (sem acento, minusculas) por `_normaliza`.
MAPA_MOVIMENTO: dict[str, TipoMovimento] = {
    "dividendo": TipoMovimento.DIVIDENDO,
    "juros sobre capital proprio": TipoMovimento.JCP,
    "jcp": TipoMovimento.JCP,
    "rendimento": TipoMovimento.RENDIMENTO,
    "amortizacao": TipoMovimento.AMORTIZACAO,
    "juros": TipoMovimento.JUROS,
    "pagamento de juros": TipoMovimento.JUROS,
    "compra": TipoMovimento.COMPRA,
    "venda": TipoMovimento.VENDA,
    "transferencia - liquidacao": None,  # o sinal da coluna Entrada/Saida decide
    "transferencia": None,  # a coluna Entrada/Saida decide a direcao
    "transferencia de custodia": None,
    "bonificacao em ativos": TipoMovimento.BONIFICACAO,
    "desdobro": TipoMovimento.DESDOBRAMENTO,
    "desdobramento": TipoMovimento.DESDOBRAMENTO,
    "grupamento": TipoMovimento.GRUPAMENTO,
    "atualizacao": TipoMovimento.OUTRO,
    "fracao em ativos": TipoMovimento.FRACAO,
    "leilao de fracao": TipoMovimento.FRACAO,
    "direito de subscricao": TipoMovimento.SUBSCRICAO,
    "direitos de subscricao": TipoMovimento.SUBSCRICAO,
    "recibo de subscricao": TipoMovimento.SUBSCRICAO,
    "cessao de direitos": TipoMovimento.SUBSCRICAO,
    "cessao de direitos - solicitada": TipoMovimento.SUBSCRICAO,
    "resgate": TipoMovimento.RESGATE,
    "aplicacao": TipoMovimento.APLICACAO,
    "cobranca de taxa semestral": TipoMovimento.TAXA,
    "incorporacao": TipoMovimento.OUTRO,
    "emprestimo": TipoMovimento.OUTRO,
}

#: Tipos que o parser entende como provento mesmo sem quantidade preenchida.
_SEM_QUANTIDADE = {
    TipoMovimento.DIVIDENDO,
    TipoMovimento.JCP,
    TipoMovimento.RENDIMENTO,
    TipoMovimento.AMORTIZACAO,
    TipoMovimento.JUROS,
}


def _normaliza(texto: Any) -> str:
    if texto is None:
        return ""
    s = str(texto).strip().lower()
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c))


def _decimal(valor: Any) -> Decimal:
    """Converte celula da B3 para Decimal.

    Aceita numero puro do openpyxl e string no formato brasileiro, incluindo o
    traco que a B3 usa para campo vazio.
    """
    if valor is None:
        return Decimal(0)
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))
    s = str(valor).strip()
    if not s or s in {"-", "--"}:
        return Decimal(0)
    s = s.replace("R$", "").replace("US$", "").strip()
    # Formato brasileiro: ponto separa milhar, virgula separa decimal.
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return Decimal(s or "0")
    except InvalidOperation:
        return Decimal(0)


def _data(valor: Any) -> Optional[date]:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    s = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(s, formato).date()
        except ValueError:
            continue
    return None


def separa_produto(produto: str) -> tuple[str, str]:
    """Quebra "PETR4 - PETROLEO BRASILEIRO S.A." em ticker e nome.

    Tesouro Direto e renda fixa nao seguem esse padrao e vem com o nome
    completo, sem hifen. Nesse caso o proprio nome vira o identificador.
    """
    texto = str(produto or "").strip()
    if " - " in texto:
        ticker, _, nome = texto.partition(" - ")
        ticker = ticker.strip()
        # O ticker real nao tem espaco. "Tesouro Selic 2029 - ..." cai fora.
        if " " not in ticker and len(ticker) <= 12:
            return ticker.upper(), nome.strip()
    return texto.upper(), texto


@dataclass
class ResultadoImport:
    novos: int = 0
    duplicados: int = 0
    ignorados: int = 0
    linhas_lidas: int = 0
    avisos: list[str] = field(default_factory=list)

    def resumo(self) -> str:
        return (
            f"{self.linhas_lidas} linhas lidas, {self.novos} novas, "
            f"{self.duplicados} ja existentes, {self.ignorados} ignoradas"
        )


def _localiza_cabecalho(linhas: list[tuple]) -> int:
    """Acha a linha do cabecalho.

    A B3 as vezes exporta com linhas de titulo antes da tabela, entao procurar
    pela celula "Entrada/Saida" e mais seguro do que assumir a primeira linha.
    """
    for indice, linha in enumerate(linhas[:15]):
        celulas = {_normaliza(c) for c in linha}
        if "entrada/saida" in celulas or ("data" in celulas and "movimentacao" in celulas):
            return indice
    return 0


def importar(
    sessao: Session,
    arquivo: BinaryIO | str | Path,
    *,
    nome_conta_padrao: str = "B3",
) -> ResultadoImport:
    """Le o Excel e grava as movimentacoes que ainda nao existem."""
    resultado = ResultadoImport()
    planilha = load_workbook(arquivo, read_only=True, data_only=True)
    aba = planilha[planilha.sheetnames[0]]
    linhas = list(aba.iter_rows(values_only=True))
    if not linhas:
        resultado.avisos.append("Planilha vazia.")
        return resultado

    inicio = _localiza_cabecalho(linhas)
    cabecalho = [_normaliza(c) for c in linhas[inicio]]

    def coluna(*nomes: str) -> Optional[int]:
        for nome in nomes:
            if nome in cabecalho:
                return cabecalho.index(nome)
        return None

    c_sentido = coluna("entrada/saida", "entrada / saida")
    c_data = coluna("data")
    c_mov = coluna("movimentacao")
    c_produto = coluna("produto")
    c_inst = coluna("instituicao")
    c_qtd = coluna("quantidade")
    c_preco = coluna("preco unitario")
    c_valor = coluna("valor da operacao", "valor")

    faltando = [n for n, c in [("Data", c_data), ("Movimentacao", c_mov), ("Produto", c_produto)] if c is None]
    if faltando:
        resultado.avisos.append(
            f"Colunas obrigatorias ausentes: {', '.join(faltando)}. "
            "Confira se o arquivo e mesmo o de Movimentacao da Area do Investidor."
        )
        return resultado

    for linha in linhas[inicio + 1 :]:
        if not any(linha):
            continue
        resultado.linhas_lidas += 1

        data_mov = _data(linha[c_data])
        if data_mov is None:
            resultado.ignorados += 1
            continue

        descricao = str(linha[c_mov] or "").strip()
        chave = _normaliza(descricao)
        sentido = _normaliza(linha[c_sentido]) if c_sentido is not None else ""
        entrada = sentido.startswith("credito") or sentido.startswith("entrada")

        tipo = MAPA_MOVIMENTO.get(chave, "_ausente")
        if tipo == "_ausente":
            # Chave desconhecida: tenta casar por prefixo antes de desistir.
            tipo = next(
                (v for k, v in MAPA_MOVIMENTO.items() if k and chave.startswith(k)),
                None,
            )
            if tipo is None and chave:
                resultado.avisos.append(f'Movimentacao nao mapeada: "{descricao}"')
        if tipo is None:
            # "Transferencia - Liquidacao" e como a B3 nomeia a liquidacao de
            # uma ordem, entao vira compra ou venda. "Transferencia" sozinha e
            # troca de custodia, que move quantidade sem ser negocio.
            if chave.startswith("transferencia - liquidacao"):
                tipo = TipoMovimento.COMPRA if entrada else TipoMovimento.VENDA
            else:
                tipo = (
                    TipoMovimento.TRANSFERENCIA_ENTRADA
                    if entrada
                    else TipoMovimento.TRANSFERENCIA_SAIDA
                )

        ticker, nome = separa_produto(linha[c_produto])
        if not ticker:
            resultado.ignorados += 1
            continue

        quantidade = _decimal(linha[c_qtd]) if c_qtd is not None else Decimal(0)
        preco = _decimal(linha[c_preco]) if c_preco is not None else Decimal(0)
        valor = _decimal(linha[c_valor]) if c_valor is not None else Decimal(0)

        if valor == 0 and quantidade and preco:
            valor = (quantidade * preco).quantize(Decimal("0.01"))

        if tipo not in _SEM_QUANTIDADE and quantidade == 0 and valor == 0:
            resultado.ignorados += 1
            continue

        nome_conta = str(linha[c_inst] or "").strip() if c_inst is not None else ""
        conta = obter_ou_criar_conta(
            sessao,
            nome_conta or nome_conta_padrao,
            instituicao=Instituicao.B3,
        )
        ativo = obter_ou_criar_ativo(sessao, ticker, nome=nome)

        mov = Movimentacao(
            data=data_mov,
            conta_id=conta.id,
            ativo_id=ativo.id,
            tipo=tipo,
            quantidade=abs(quantidade),
            preco_unitario=preco or None,
            valor_bruto=abs(valor),
            valor_liquido=abs(valor),
            moeda="BRL",
            fonte=Fonte.B3_MOVIMENTACAO,
            descricao=descricao,
        )
        if registrar_movimentacao(sessao, mov):
            resultado.novos += 1
        else:
            resultado.duplicados += 1

    sessao.commit()
    # Avisos repetidos poluem: um por tipo de movimentacao desconhecida basta.
    resultado.avisos = sorted(set(resultado.avisos))
    return resultado
